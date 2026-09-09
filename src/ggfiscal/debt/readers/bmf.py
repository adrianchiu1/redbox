"""Readers for family 'bmf': BMF snapshots -> tidy frames (DEBT_KICKOFF.md §6.3).

Three shapes:

* **Datenportal workbook** (``datenportal``). One xlsx, four wide sheets —
  ``rpgSchuldenstand`` (stock, month-ends from 1995-12-31),
  ``rpgBruttokreditaufnahme Gesamt``, ``rpgTilgungen``, ``rpgZinsen Gesamt``
  (from 1996 M01). Rows are a *tree*: the hierarchy is carried by the cell
  indentation of column A, never by the label text, and is reconstructed here
  into a ``series_path``. Values are whole euro. **Both flow sheets and the
  interest sheet are cumulative within the calendar year** (1996 M01 = January,
  1996 M12 = the full year, 1997 M01 starts again), and Tilgungen and Zinsen
  are signed negative (payments out); the stock sheet is a level.
* **Datenportal CSVs** (``datenportal_csv``). The same three series published
  one per file, UTF-16LE, tab-separated, wide, with the section markers kept
  but the indentation lost — so the CSVs give labels and rows, the workbook
  gives the tree. The CSVs carry a few instrument rows the workbook's
  Instrumentenarten block does not show (the inflation-linked and green
  sub-lines).
* **Kreditaufnahmebericht PDFs** (``kreditaufnahmebericht_text`` and
  ``kreditaufnahmebericht_annex``). Text-layer extraction with pypdf; annex 4.5
  (Verzinsung, incl. the "nachrichtlich: Verzinsung gemäß Epl. 32, Kapitel
  3205" block) and annex 4.10 (Abrechnung des Kreditfinanzierungsplans, Soll vs
  Ist) are parsed from the extracted line layout. What parses per edition is
  recorded in ``reports/debt_sources/bmf.md``; nothing is ever hand-keyed.
"""

from __future__ import annotations

import re
from functools import lru_cache

import pandas as pd

from ggfiscal.debt.readers import snap_path

DATENPORTAL_XLSX = "kredit_brutto_tilgung_zinsen_xlsx"
SHEETS = ("rpgSchuldenstand", "rpgBruttokreditaufnahme Gesamt",
          "rpgTilgungen", "rpgZinsen Gesamt")
#: Datenportal CSV part -> the workbook sheet carrying the same series.
CSV_SHEET = {
    "bruttokreditaufnahme_csv": "rpgBruttokreditaufnahme Gesamt",
    "tilgungen_csv": "rpgTilgungen",
    "kreditbestand_csv": "rpgSchuldenstand",
}

#: Rows above the first "Gliederung nach ..." marker are the sheet totals; they
#: belong to no breakdown, so they get their own section name.
TOTAL_SECTION = "Gesamt"
SECTION_RE = re.compile(r"^Gliederung nach\s+(.+?)\s*$")
NACHRICHTLICH_RE = re.compile(r"^Nachrichtlich\b")
PERIOD_DMY_RE = re.compile(r"^(\d{2})\.(\d{2})\.(\d{4})$")
PERIOD_MONTH_RE = re.compile(r"^(\d{4})\s*M(\d{2})$")


def _period(label: object) -> pd.Timestamp | None:
    """'31.12.1995' or '1996 M01' -> the month-end Timestamp."""
    text = str(label).strip()
    m = PERIOD_DMY_RE.match(text)
    if m:
        day, month, year = (int(g) for g in m.groups())
        return pd.Timestamp(year=year, month=month, day=day)
    m = PERIOD_MONTH_RE.match(text)
    if m:
        year, month = int(m.group(1)), int(m.group(2))
        return pd.Timestamp(year=year, month=month, day=1) + pd.offsets.MonthEnd(0)
    return None


# ---------------------------------------------------------------- Datenportal xlsx

@lru_cache(maxsize=None)
def _workbook():
    import openpyxl

    # data_only=True for cached values; not read_only, because the row
    # hierarchy lives in each cell's alignment.indent.
    return openpyxl.load_workbook(snap_path("BMF_DATENPORTAL", DATENPORTAL_XLSX),
                                  data_only=True)


@lru_cache(maxsize=None)
def datenportal(sheet: str) -> pd.DataFrame:
    """One Datenportal sheet -> long frame.

    Columns: ``series_path`` (section and ancestors joined with " / "),
    ``series_label`` (the leaf), ``section`` (Gesamt | Verwendung |
    Instrumentenarten | Restlaufzeiten | Währungen | Nachrichtlich),
    ``level`` (indentation depth), ``period`` (month-end), ``value_eur``,
    ``value_eur_mn``.
    """
    if sheet not in SHEETS:
        raise KeyError(f"unknown Datenportal sheet {sheet!r}; expected one of {SHEETS}")
    ws = _workbook()[sheet]

    periods: dict[int, pd.Timestamp] = {}
    header_row = 0
    for i in range(1, ws.max_row + 1):
        if ws.cell(i, 1).value not in (None, "") :
            continue
        cand = {j: _period(ws.cell(i, j).value)
                for j in range(2, ws.max_column + 1)
                if ws.cell(i, j).value not in (None, "")}
        if cand and all(v is not None for v in cand.values()):
            periods, header_row = {j: v for j, v in cand.items()}, i
            break
    if not periods:
        raise ValueError(f"no period header row found in sheet {sheet!r}")

    section = TOTAL_SECTION
    stack: dict[int, str] = {}
    records: list[dict] = []
    for i in range(header_row + 1, ws.max_row + 1):
        cell = ws.cell(i, 1)
        label = "" if cell.value is None else str(cell.value).strip()
        if not label:
            continue
        m = SECTION_RE.match(label)
        if m:
            section, stack = m.group(1).strip(), {}
            continue
        if NACHRICHTLICH_RE.match(label):
            section, stack = "Nachrichtlich", {}
            continue
        values = {}
        for j, period in periods.items():
            v = pd.to_numeric(ws.cell(i, j).value, errors="coerce")
            if pd.notna(v):
                values[period] = float(v)
        if not values:
            continue  # footnote or spacer row
        level = int(cell.alignment.indent or 0)
        stack = {d: lab for d, lab in stack.items() if d < level}
        stack[level] = label
        path = " / ".join([section] + [stack[d] for d in sorted(stack)])
        records += [{"series_path": path, "series_label": label,
                     "section": section, "level": level, "period": period,
                     "value_eur": value, "value_eur_mn": value / 1e6}
                    for period, value in values.items()]
    return (pd.DataFrame.from_records(records)
            .sort_values(["series_path", "period"], kind="stable")
            .reset_index(drop=True))


def _instrument_block(sheet: str, extras: tuple[str, ...] = ()) -> pd.DataFrame:
    df = datenportal(sheet)
    keep = df["section"] == "Instrumentenarten"
    for prefix in extras:
        keep |= df["series_path"].str.startswith(f"Nachrichtlich / {prefix}")
    return (df.loc[keep, ["series_path", "series_label", "section", "period",
                          "value_eur_mn"]]
              .reset_index(drop=True))


def bund_stock_by_instrument() -> pd.DataFrame:
    """Kreditbestand by instrument, EUR mn, month-ends from 1995-12-31.

    The Instrumentenarten tree plus the Nachrichtlich Umlaufvolumen tree (the
    same instruments measured as amount in issue) and the Eigenbestände lines
    (DD11/§14: allotted is not issued — market-hands = Umlaufvolumen +
    Eigenbestände, the Eigenbestände being carried negative)."""
    return _instrument_block("rpgSchuldenstand", ("Umlaufvolumen", "Eigenbestände"))


def bund_gross_issuance_by_instrument() -> pd.DataFrame:
    """Bruttokreditaufnahme by instrument, EUR mn, cumulative within each year."""
    return _instrument_block("rpgBruttokreditaufnahme Gesamt")


def bund_redemptions_by_instrument() -> pd.DataFrame:
    """Tilgungen by instrument, EUR mn, cumulative within each year, negative."""
    return _instrument_block("rpgTilgungen")


def bund_interest_by_instrument() -> pd.DataFrame:
    """Verzinsung by instrument, EUR mn, cumulative within each year, negative
    (the sheet is the balance of interest paid and received, so an interest
    *cost* carries a minus sign)."""
    return _instrument_block("rpgZinsen Gesamt")


# ---------------------------------------------------------------- Datenportal CSV

@lru_cache(maxsize=None)
def datenportal_csv(part: str) -> pd.DataFrame:
    """One Datenportal CSV (UTF-16LE, tab-separated, wide) -> long frame.

    Columns: ``row_index`` (position in the file, the only way to tell repeated
    labels apart once the indentation is gone), ``series_label``, ``section``,
    ``period``, ``value_eur``, ``value_eur_mn``.
    """
    if part not in CSV_SHEET:
        raise KeyError(f"unknown Datenportal CSV part {part!r}; "
                       f"expected one of {tuple(CSV_SHEET)}")
    text = snap_path("BMF_DATENPORTAL", part).read_bytes().decode("utf-16")
    lines = [line.split("\t") for line in text.splitlines() if line.strip()]
    head = lines[0]
    periods = {j: _period(v) for j, v in enumerate(head) if j and _period(v)}
    if not periods:
        raise ValueError(f"no period header found in CSV part {part!r}")

    section = TOTAL_SECTION
    records: list[dict] = []
    for row_index, fields in enumerate(lines[1:]):
        label = fields[0].strip()
        if not label:
            continue
        m = SECTION_RE.match(label)
        if m:
            section = m.group(1).strip()
            continue
        if NACHRICHTLICH_RE.match(label):
            section = "Nachrichtlich"
            continue
        for j, period in periods.items():
            if j >= len(fields):
                continue
            value = pd.to_numeric(fields[j].strip().replace(" ", ""), errors="coerce")
            if pd.notna(value):
                records.append({"row_index": row_index, "series_label": label,
                                "section": section, "period": period,
                                "value_eur": float(value),
                                "value_eur_mn": float(value) / 1e6})
    return (pd.DataFrame.from_records(records)
            .sort_values(["row_index", "period"], kind="stable")
            .reset_index(drop=True))


# ------------------------------------------------------- Kreditaufnahmebericht PDF

#: A German thousands-grouped amount, with or without decimals.
_NUM = r"-?\d{1,3}(?:\.\d{3})*(?:,\d+)?"
#: One table cell: an amount, or "-" for "no value in this period".
_CELL_RE = re.compile(rf"(?:(?<=\s)|^)(?:{_NUM}|-)(?=\s|$)")
_YEAR_ROW_RE = re.compile(r"^(?:(?:19|20)\d{2}\s+){1,}(?:19|20)\d{2}$")
_RUNNING_HEAD_RE = re.compile(r"^(Bericht des Bundesministeriums|Berichtsperiode|"
                              r"in Mio\.|in €|Bezeichnung|Noch\s|\*|Soll\s|\d{1,3}$)")
_ITEM_RE = re.compile(r"^(\d+(?:\.\d+)*)\s+(.*)$")
#: A group heading in annex 4.10 ("1 Einnahmen"): a top-level item number with
#: no value columns of its own.
_GROUP_RE = re.compile(r"^(\d+)\s+(\S.*)$")


def _join_label(parts: list[str]) -> str:
    """Join the fragments of a label wrapped across PDF lines, repairing the
    typesetter's soft hyphens ("Haushaltsausga- ben" -> "Haushaltsausgaben")."""
    out = ""
    for part in (p.strip() for p in parts if p and p.strip()):
        if out.endswith("-") and part[:1].islower():
            out = out[:-1] + part
        elif out:
            out = f"{out} {part}"
        else:
            out = part
    return re.sub(r"\s+", " ", out).strip()


def _to_float(token: str) -> float | None:
    token = token.strip()
    if token in ("-", "–", ""):
        return None
    return float(token.replace(".", "").replace(",", "."))


def _trailing_cells(line: str) -> tuple[str, list[str]]:
    """Split a table line into (label text, trailing numeric cells)."""
    cells: list[str] = []
    rest = line.rstrip()
    while True:
        m = None
        for m in _CELL_RE.finditer(rest):
            pass
        if m is None or m.end() != len(rest):
            break
        cells.insert(0, m.group(0))
        rest = rest[:m.start()].rstrip()
    return rest.strip(), cells


@lru_cache(maxsize=None)
def kreditaufnahmebericht_pages(year: int) -> tuple[str, ...]:
    """The PDF's text layer, one string per page (pypdf)."""
    from pypdf import PdfReader

    reader = PdfReader(snap_path("BMF_KREDITAUFNAHMEBERICHT", str(year)))
    return tuple((page.extract_text() or "") for page in reader.pages)


def kreditaufnahmebericht_text(year: int) -> str:
    """The whole edition's text layer, pages separated by a form feed."""
    return "\f".join(kreditaufnahmebericht_pages(year))


#: How to find an annex in an edition. The *number* of an annex moves between
#: editions (the Verzinsung table is 4.10 in 2019, 4.6 in 2020, 4.5 from 2022;
#: the Abrechnung is 4.11 in 2020-21 and 4.10 from 2022), so the annex is found
#: by its title and its unit line, and the number is read off the heading that
#: matched. The keys are the 2025 numbering, as used by DEBT_KICKOFF.md §6.3.
_ANNEX_SPECS = {
    "4.5": (re.compile(r"^(?:Anhang\s*)?(\d+\.\d+)\s*:?\s+"
                       r"(?:nachrichtlich:\s*)?Verzinsung\b", re.M), "in Mio."),
    "4.10": (re.compile(r"^(?:Anhang\s*)?(\d+\.\d+)\s*:?\s+"
                        r"Abrechnung des Kreditfinanzierungsplans", re.M), "in €"),
}


def _annex_pages(year: int, annex: str) -> list[str]:
    """The pages of one annex, in order: the page whose heading matches the
    annex title, plus the continuation pages ("Noch 4.5 ...", "4.11:
    Fortsetzung") that repeat its number. Requiring the unit line as well drops
    the table of contents, which repeats every title."""
    title_re, unit = _ANNEX_SPECS[annex]
    pages = kreditaufnahmebericht_pages(year)
    first = None
    for i, page in enumerate(pages):
        m = title_re.search(page)
        # Chapter 4 is the annex in every edition that carries these tables;
        # the same words appear in chapter 1-3 body headings, which are not it.
        if m and m.group(1).startswith("4.") and unit in page:
            first, number = i, m.group(1)
            break
    if first is None:
        return []
    out = [pages[first]]
    number_re = re.compile(rf"(?<!\d){re.escape(number)}(?!\d)")
    other_title_re = re.compile(r"^(?:Anhang\s*)?(4\.\d+)\s*:?\s+\S", re.M)
    item_line_re = re.compile(r"^\s*\d+\.\d+(?:\.\d+)?\s+\S", re.M)
    for page in pages[first + 1:]:
        # a page that opens another annex (a different 4.x number) ends this one
        titles = {m.group(1) for m in other_title_re.finditer(page)} - {number}
        if titles and not number_re.search(page):
            break
        if unit in page and (number_re.search(page) or "Fortsetzung" in page or "Noch" in page
                             or item_line_re.search(page)):
            out.append(page)
        else:
            break
    return out


def _page_rows_45(page: str) -> list[dict]:
    """One annex-4.5 page -> row dicts (section, label, years, cells).

    Labels appear only on the first page of each table; the continuation pages
    (the later year columns) carry the value columns alone, in the same row
    order, so their labels are filled in by :func:`_parse_45`.
    """
    rows: list[dict] = []
    years: list[int] = []
    section = ""
    buffer: list[str] = []
    for raw in page.splitlines():
        line = raw.strip()
        if not line:
            continue
        if _YEAR_ROW_RE.match(line):
            years, buffer = [int(t) for t in line.split()], []
            continue
        m = SECTION_RE.match(line)
        if m:
            section, buffer = m.group(1).strip(), []
            continue
        if _RUNNING_HEAD_RE.match(line):
            buffer = []
            continue
        label, cells = _trailing_cells(line)
        if not years or len(cells) != len(years):
            buffer.append(line)
            continue
        rows.append({"section": section, "label": _join_label([*buffer, label]),
                     "years": years, "cells": cells})
        buffer = []
    return rows


def _parse_45(year: int) -> pd.DataFrame:
    """Annex 4.5 — Verzinsung des Kreditbestands, and the nachrichtlich Epl. 32
    Kapitel 3205 block. Year columns, EUR mn.

    The table is broken across pages by year block; only the first page of each
    block carries the row labels in its text layer. A label-less page is
    accepted only when its row count matches the labelled page it continues —
    otherwise its rows are dropped rather than guessed at (V-rule: never key a
    number to a label we cannot demonstrate).
    """
    records: list[dict] = []
    carried: dict[str, list[tuple[str, str]]] = {}
    for page in _annex_pages(year, "4.5"):
        table = ("kap3205" if re.search(r"nachrichtlich.*Epl\.\s*32", page, re.I)
                 else "verzinsung")
        rows = _page_rows_45(page)
        if not rows:
            continue
        if any(r["label"] for r in rows):
            carried[table] = [(r["section"], r["label"]) for r in rows]
        else:
            labels = carried.get(table, [])
            if len(labels) != len(rows):
                continue
            for row, (section, label) in zip(rows, labels):
                row["section"], row["label"] = section, label
        for row in rows:
            if not row["label"]:
                continue
            for y, cell in zip(row["years"], row["cells"]):
                value = _to_float(cell)
                if value is not None:
                    records.append({"annex": "4.5", "table": table,
                                    "section": row["section"],
                                    "series_label": row["label"],
                                    "year": y, "value_eur_mn": value})
    return pd.DataFrame.from_records(records)


def _parse_410(year: int) -> pd.DataFrame:
    """Annex 4.10 — Abrechnung des Kreditfinanzierungsplans: item number, group
    heading, label, and the three columns Soll / Ist / Abweichung, in euro.

    Unnumbered running subtotals (the bare amounts the table sets between
    items 3.2/3.3 and 3.3/3.4) carry no label in the text layer and are
    dropped; every labelled line, including the "Einnahmen", "Ausgaben" and
    "Nettokreditaufnahme" totals, is kept.
    """
    records: list[dict] = []
    group = ""          # carried across pages: a continuation page repeats no heading
    for page in _annex_pages(year, "4.10"):
        buffer: list[str] = []
        for raw in page.splitlines():
            line = raw.strip()
            if not line:
                continue
            if _RUNNING_HEAD_RE.match(line) and not _ITEM_RE.match(line):
                buffer = []
                continue
            label, cells = _trailing_cells(line)
            if len(cells) != 3:
                m = _GROUP_RE.match(line)
                if m:                      # "1 Einnahmen": a heading, not a label
                    group, buffer = m.group(2).strip(), []
                else:
                    buffer.append(line)
                continue
            full_label = _join_label([*buffer, label])
            buffer = []
            m = _ITEM_RE.match(full_label)
            item, text = (m.group(1), m.group(2)) if m else ("", full_label)
            if not text:
                continue
            soll, ist, abw = (_to_float(c) for c in cells)
            records.append({"annex": "4.10", "item": item, "group": group,
                            "label": text, "soll_eur": soll, "ist_eur": ist,
                            "abweichung_eur": abw})
    return pd.DataFrame.from_records(records)


ANNEX_PARSERS = {"4.5": _parse_45, "4.10": _parse_410}


def kreditaufnahmebericht_annex(year: int, annex: str) -> pd.DataFrame:
    """Parse one Kreditaufnahmebericht annex out of the PDF's text layer.

    ``4.5`` returns (annex, table, section, series_label, year, value_eur_mn)
    where ``table`` is ``verzinsung`` or ``kap3205``; ``4.10`` returns
    (annex, item, group, label, soll_eur, ist_eur, abweichung_eur). An edition whose
    text layer does not carry the annex returns an empty frame — the caller
    must check, and ``reports/debt_sources/bmf.md`` records which editions
    parse.
    """
    if annex not in ANNEX_PARSERS:
        raise KeyError(f"no parser for annex {annex!r}; have {tuple(ANNEX_PARSERS)}")
    return ANNEX_PARSERS[annex](year)


# ------------------------------------------------ annual totals for the chains

def bund_interest_annual(edition: int = 2025) -> pd.Series:
    """Step-A interest total for DEU (DEBT_KICKOFF.md §6.3): the
    Kreditaufnahmebericht annex 4.5 'Insgesamt' row — Verzinsung des Bundes
    (net of interest income; excludes swaps and cash management, §14 DEU),
    1996 onward in the 2025 edition. Sign flipped to positive expenditure.
    Cross-checked in the notes against the Datenportal Zinsen sheet's
    December (year-to-date) total."""
    a = kreditaufnahmebericht_annex(edition, "4.5")
    tot = a[(a["series_label"] == "Insgesamt") & (a["table"] == "verzinsung")]
    s = (-tot.set_index("year")["value_eur_mn"]).sort_index().astype(float)
    try:
        z = datenportal("rpgZinsen Gesamt")
        z = z[(z["level"] == 0) & (z["period"].dt.month == 12)]
        z = z[z["series_path"].str.contains("Finanzierung Bundeshaushalt")]
        dp = (-z.set_index(z["period"].dt.year)["value_eur_mn"]).astype(float)
        diff = (s.reindex(dp.index) - dp).abs().max()
        s.attrs["note"] = (f"Kreditaufnahmebericht {edition} annex 4.5 Insgesamt (Verzinsung, net of interest "
                           f"income; excl. swaps); max |diff| vs Datenportal Zinsen Dec YTD = {diff:,.1f} EUR mn")
    except Exception as e:      # cross-check is informative only
        s.attrs["note"] = f"Kreditaufnahmebericht {edition} annex 4.5 Insgesamt; Datenportal cross-check failed: {e}"
    return s


def bund_net_borrowing_annual() -> pd.Series:
    """Step-A financing total for DEU: Nettokreditaufnahme (Ist) from each
    edition's annex 4.10 'Abrechnung des Kreditfinanzierungsplans', in EUR
    mn, for every edition whose annex parses (2020 onward; earlier
    editions' annexes are not machine-readable in the text layer)."""
    from ggfiscal.standardise.readers import latest_snapshots
    out = {}
    for (sid, part) in latest_snapshots():
        if sid != "BMF_KREDITAUFNAHMEBERICHT" or not part.isdigit():
            continue
        try:
            t = kreditaufnahmebericht_annex(int(part), "4.10")
        except Exception:
            continue
        if t.empty:
            continue
        r = t[t["label"].str.strip().str.lower().str.startswith("nettokreditaufnahme")]
        if r.empty:
            continue
        out[int(part)] = float(r["ist_eur"].iloc[0]) / 1e6
    s = pd.Series(out, dtype=float).sort_index()
    s.attrs["note"] = "Kreditaufnahmebericht annex 4.10 Nettokreditaufnahme (Ist), edition = year"
    return s
